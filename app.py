"""
app.py — Informe Automático BCE
Interfaz Streamlit para generar informes de balance en formato Word.
"""

import streamlit as st
import os
import json
import base64
import tempfile
import time
from pathlib import Path
import urllib.request

# ── Configuración de página ──────────────────────────────────────────────────
st.set_page_config(
    page_title="Informe Automático BCE",
    page_icon="🏦",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# ── Estilos ──────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    /* Fondo y fuente general */
    .stApp { background-color: #F7F8FA; }
    html, body, [class*="css"] { font-family: 'Segoe UI', sans-serif; }

    /* Header institucional */
    .header-bna {
        background: linear-gradient(135deg, #003087 0%, #0052CC 100%);
        color: white;
        padding: 1.4rem 2rem;
        border-radius: 10px;
        margin-bottom: 1.8rem;
        display: flex;
        align-items: center;
        gap: 1rem;
    }
    .header-bna h1 { margin: 0; font-size: 1.45rem; font-weight: 700; letter-spacing: 0.01em; }
    .header-bna p  { margin: 0.2rem 0 0; font-size: 0.82rem; opacity: 0.82; }

    /* Tarjeta de carga */
    .upload-card {
        background: white;
        border: 2px dashed #C5D3E8;
        border-radius: 10px;
        padding: 1.4rem 1.6rem;
        margin-bottom: 1rem;
        transition: border-color 0.2s;
    }
    .upload-card:hover { border-color: #0052CC; }
    .upload-label {
        font-size: 0.78rem;
        font-weight: 600;
        text-transform: uppercase;
        letter-spacing: 0.07em;
        color: #5A6A85;
        margin-bottom: 0.4rem;
    }

    /* Botón principal */
    .stButton > button {
        background: linear-gradient(135deg, #003087 0%, #0052CC 100%);
        color: white !important;
        border: none;
        border-radius: 8px;
        padding: 0.65rem 2rem;
        font-size: 0.95rem;
        font-weight: 600;
        width: 100%;
        cursor: pointer;
        transition: opacity 0.2s;
    }
    .stButton > button:hover { opacity: 0.88; }
    .stButton > button:disabled { opacity: 0.45; cursor: not-allowed; }

    /* Pasos del proceso */
    .step-box {
        background: white;
        border-left: 4px solid #0052CC;
        border-radius: 0 8px 8px 0;
        padding: 0.7rem 1rem;
        margin: 0.4rem 0;
        font-size: 0.88rem;
        color: #1A2B4A;
    }
    .step-box.done  { border-left-color: #00875A; color: #00875A; }
    .step-box.error { border-left-color: #DE350B; color: #DE350B; }

    /* Resultado final */
    .result-card {
        background: #E3FCEF;
        border: 1.5px solid #00875A;
        border-radius: 10px;
        padding: 1.2rem 1.6rem;
        text-align: center;
        margin-top: 1rem;
    }
    .result-card h3 { color: #00875A; margin: 0 0 0.3rem; font-size: 1.1rem; }
    .result-card p  { color: #1A2B4A; margin: 0; font-size: 0.85rem; }

    /* Info box */
    .info-box {
        background: #EBF2FF;
        border: 1px solid #B3CEFF;
        border-radius: 8px;
        padding: 0.8rem 1rem;
        font-size: 0.83rem;
        color: #1A2B4A;
        margin-bottom: 1rem;
    }

    /* Ocultar elementos Streamlit */
    #MainMenu, footer, header { visibility: hidden; }
    .block-container { padding-top: 1.5rem; max-width: 720px; }
</style>
""", unsafe_allow_html=True)

# ── Login simple ─────────────────────────────────────────────────────────────
USUARIOS = {
    "analista1": "bna2024",
    "analista2": "bna2024",
    "supervisor": "bna2025",
}

def login():
    st.markdown("""
    <div class="header-bna">
        <div>🏦</div>
        <div>
            <h1>Informe Automático BCE</h1>
            <p>Banco de la Nación Argentina — Acceso restringido</p>
        </div>
    </div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("#### Iniciar sesión")
        usuario = st.text_input("Usuario", key="login_user")
        clave   = st.text_input("Contraseña", type="password", key="login_pass")
        if st.button("Ingresar"):
            if USUARIOS.get(usuario) == clave:
                st.session_state["autenticado"] = True
                st.session_state["usuario"]     = usuario
                st.rerun()
            else:
                st.error("Usuario o contraseña incorrectos.")

# ── Helpers API ───────────────────────────────────────────────────────────────
API_URL = "https://api.anthropic.com/v1/messages"
MODEL   = "claude-sonnet-4-20250514"

def pdf_a_base64(bytes_pdf: bytes) -> str:
    return base64.standard_b64encode(bytes_pdf).decode("utf-8")

def llamar_api(prompt: str, pdf_b64: str, max_tokens: int, api_key: str) -> dict:
    payload = json.dumps({
        "model": MODEL,
        "max_tokens": max_tokens,
        "messages": [{
            "role": "user",
            "content": [
                {"type": "document",
                 "source": {"type": "base64", "media_type": "application/pdf", "data": pdf_b64}},
                {"type": "text", "text": prompt}
            ]
        }]
    }).encode("utf-8")

    req = urllib.request.Request(
        API_URL, data=payload,
        headers={"Content-Type": "application/json",
                 "x-api-key": api_key,
                 "anthropic-version": "2023-06-01"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        r = json.loads(resp.read().decode("utf-8"))
    texto = r["content"][0]["text"].strip()
    if texto.startswith("```"):
        texto = texto.split("\n", 1)[1].rsplit("```", 1)[0]
    return json.loads(texto)

def llamar_api_texto(prompt: str, api_key: str, max_tokens: int = 3000) -> dict:
    payload = json.dumps({
        "model": MODEL,
        "max_tokens": max_tokens,
        "messages": [{"role": "user", "content": prompt}]
    }).encode("utf-8")

    req = urllib.request.Request(
        API_URL, data=payload,
        headers={"Content-Type": "application/json",
                 "x-api-key": api_key,
                 "anthropic-version": "2023-06-01"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        r = json.loads(resp.read().decode("utf-8"))
    texto = r["content"][0]["text"].strip()
    if texto.startswith("```"):
        texto = texto.split("\n", 1)[1].rsplit("```", 1)[0]
    return json.loads(texto)

# ── Extracción de datos ───────────────────────────────────────────────────────
PROMPT_BALANCE = """Sos un analista contable argentino. Analizá estos estados contables y extraé TODOS los datos numéricos y textuales listados.
Devolvé ÚNICAMENTE un objeto JSON válido, sin texto adicional, sin markdown, sin backticks.

{
  "razon_social": "",
  "cuit": "",
  "domicilio": "",
  "actividad": "",
  "fecha_cierre": "",
  "fecha_cierre_anterior": "",
  "nro_ejercicio": 0,
  "consejo_protocolizante": "",
  "nombre_auditor": "",
  "tiene_salvedades": false,
  "activo_corriente": { "caja_bancos": 0, "inversiones": 0, "creditos_por_ventas": 0, "otros_creditos": 0, "bienes_de_cambio": 0, "total": 0 },
  "activo_no_corriente": { "bienes_de_uso": 0, "inversiones_permanentes": 0, "total": 0 },
  "total_activo": 0,
  "pasivo_corriente": { "deudas_comerciales": 0, "deudas_bancarias": 0, "remuneraciones": 0, "cargas_fiscales": 0, "otras_deudas": 0, "total": 0 },
  "pasivo_no_corriente": { "deudas_bancarias": 0, "cargas_fiscales": 0, "aportes_irrevocables": 0, "total": 0 },
  "total_pasivo": 0,
  "patrimonio_neto": 0,
  "activo_corriente_anterior": { "caja_bancos": 0, "inversiones": 0, "creditos_por_ventas": 0, "otros_creditos": 0, "bienes_de_cambio": 0, "total": 0 },
  "activo_no_corriente_anterior": { "bienes_de_uso": 0, "total": 0 },
  "total_activo_anterior": 0,
  "pasivo_corriente_anterior": { "deudas_comerciales": 0, "deudas_bancarias": 0, "remuneraciones": 0, "cargas_fiscales": 0, "otras_deudas": 0, "total": 0 },
  "pasivo_no_corriente_anterior": { "total": 0 },
  "total_pasivo_anterior": 0,
  "patrimonio_neto_anterior": 0,
  "ventas_netas": 0, "otros_ingresos": 0, "costo_ventas": 0, "resultado_bruto": 0,
  "gastos_comercializacion": 0, "gastos_administracion": 0, "gastos_financieros": 0,
  "recpam": 0, "resultado_operativo": 0, "ganancia_neta": 0,
  "ventas_netas_anterior": 0, "costo_ventas_anterior": 0, "resultado_bruto_anterior": 0,
  "gastos_comercializacion_anterior": 0, "gastos_administracion_anterior": 0,
  "gastos_financieros_anterior": 0, "resultado_operativo_anterior": 0, "ganancia_neta_anterior": 0,
  "amortizaciones": 0, "compras_periodo": 0
}

IMPORTANTE:
- Importes en MILES de pesos (dividí por 1000 si están en pesos completos)
- Si un valor no figura, usá 0
- fecha_cierre formato: "DD de mes de YYYY"
- nro_ejercicio: número del ejercicio (ej: 12)
- Buscá el CUIT en el informe del auditor o constancia de legalización
"""

PROMPT_NOSIS = """Sos un analista crediticio argentino. Analizá este informe NOSIS y extraé los datos indicados.
Devolvé ÚNICAMENTE un objeto JSON válido, sin texto adicional, sin markdown, sin backticks.

{
  "razon_social": "", "cuit": "", "domicilio": "", "actividad_principal": "",
  "tipo_sociedad": "", "fecha_contrato_social": "",
  "score": 0, "resultado_cda": "",
  "endeudamiento_sf_actual": 0, "compromisos_mensuales": 0,
  "situacion_bcra": "",
  "tiene_cheques_rechazados": false, "tiene_deuda_previsional": false,
  "deuda_previsional_detalle": "",
  "srt_vigente": false, "srt_aseguradora": "",
  "cantidad_empleados": "",
  "entidades_deuda": [{"entidad": "", "monto": 0, "situacion": ""}],
  "consultas_ultimos_12_meses": 0,
  "socios": [{"nombre": "", "porcentaje": 0, "cargo": ""}]
}

IMPORTANTE:
- endeudamiento_sf_actual y compromisos_mensuales en miles de pesos
- socios: buscalos en árbol de relaciones o sección SOCIEDADES
"""

# ── Cálculo de índices ────────────────────────────────────────────────────────
def safe_div(a, b):
    try: return round(a / b, 4) if b else 0
    except: return 0

def pct(a, b):
    if not b: return None
    return round((a - b) / abs(b) * 100, 1)

def calcular_indices(b):
    ac   = b.get("activo_corriente", {})
    anc  = b.get("activo_no_corriente", {})
    pc   = b.get("pasivo_corriente", {})
    pnc  = b.get("pasivo_no_corriente", {})
    at   = b.get("total_activo", 0)
    pt   = b.get("total_pasivo", 0)
    pn   = b.get("patrimonio_neto", 0)
    ac_t = ac.get("total", 0);  pc_t = pc.get("total", 0)
    bc   = ac.get("bienes_de_cambio", 0)
    bu   = anc.get("bienes_de_uso", 0)
    cxv  = ac.get("creditos_por_ventas", 0)
    at_ant = b.get("total_activo_anterior", 0)
    pt_ant = b.get("total_pasivo_anterior", 0)
    ac_ant = b.get("activo_corriente_anterior", {}).get("total", 0)
    pc_ant = b.get("pasivo_corriente_anterior", {}).get("total", 0)
    vtas = b.get("ventas_netas", 0); cmv = b.get("costo_ventas", 0)
    rb   = b.get("resultado_bruto", 0); rop = b.get("resultado_operativo", 0)
    gn   = b.get("ganancia_neta", 0)
    gcom = b.get("gastos_comercializacion", 0)
    gadm = b.get("gastos_administracion", 0)
    gfin = b.get("gastos_financieros", 0)
    recpam = b.get("recpam", 0)
    amort = b.get("amortizaciones", 0)
    compras = b.get("compras_periodo", 0)
    vtas_ant = b.get("ventas_netas_anterior", 0)
    rb_ant   = b.get("resultado_bruto_anterior", 0)
    rop_ant  = b.get("resultado_operativo_anterior", 0)
    gn_ant   = b.get("ganancia_neta_anterior", 0)
    gcom_ant = b.get("gastos_comercializacion_anterior", 0)
    gadm_ant = b.get("gastos_administracion_anterior", 0)
    pnc_t = pnc.get("total", 0)

    lc     = safe_div(ac_t, pc_t)
    la     = safe_div(ac_t - bc, pc_t)
    lc_ant = safe_div(ac_ant, pc_ant)
    ebitda = rop + amort

    return {
        "solvencia":             round(safe_div(at, pt), 2),
        "endeudamiento":         round(safe_div(pt, pn), 2),
        "end_diferido":          round(safe_div(pnc_t, pn), 2),
        "financiamiento_propio": round(safe_div(pn, at) * 100, 1),
        "liquidez_corriente":    round(lc, 2),
        "liquidez_acida":        round(la, 2),
        "capital_trabajo":       round(ac_t - pc_t, 0),
        "rotacion_bc_dias":      round(safe_div(bc * 365, cmv), 0) if cmv else 0,
        "rotacion_cxv_dias":     round(safe_div(cxv * 365, vtas), 0) if vtas else 0,
        "margen_bruto_pct":      round(safe_div(rb, vtas) * 100, 1),
        "margen_bruto_ant_pct":  round(safe_div(rb_ant, vtas_ant) * 100, 1),
        "margen_operativo_pct":  round(safe_div(rop, vtas) * 100, 1),
        "margen_neto_pct":       round(safe_div(gn, vtas) * 100, 1),
        "margen_neto_ant_pct":   round(safe_div(gn_ant, vtas_ant) * 100, 1),
        "rent_activo_pct":       round(safe_div(gn, at) * 100, 1),
        "rent_pn_pct":           round(safe_div(gn, pn) * 100, 1),
        "ebitda":                round(ebitda, 0),
        "ebitda_pct":            round(safe_div(ebitda, vtas) * 100, 1),
        "gcom_pct":              round(safe_div(gcom, vtas) * 100, 1),
        "gadm_pct":              round(safe_div(gadm, vtas) * 100, 1),
        "gfin_pct":              round(safe_div(abs(gfin), vtas) * 100, 1),
        "recpam_pct":            round(safe_div(abs(recpam), vtas) * 100, 1),
        "promedio_mensual_vtas": round(vtas / 12, 0),
        "var_ac":   pct(ac_t, ac_ant),   "var_at":  pct(at, at_ant),
        "var_pt":   pct(pt, pt_ant),     "var_pc":  pct(pc_t, pc_ant),
        "var_vtas": pct(vtas, vtas_ant), "var_rb":  pct(rb, rb_ant),
        "var_rop":  pct(rop, rop_ant),   "var_gn":  pct(gn, gn_ant),
        "var_lc":   pct(lc, lc_ant),
        "var_bc":   pct(bc, b.get("activo_corriente_anterior",{}).get("bienes_de_cambio",0)),
        "var_cxv":  pct(cxv, b.get("activo_corriente_anterior",{}).get("creditos_por_ventas",0)),
        "var_bu":   pct(bu, b.get("activo_no_corriente_anterior",{}).get("bienes_de_uso",0)),
        "var_dc":   pct(pc.get("deudas_comerciales",0), b.get("pasivo_corriente_anterior",{}).get("deudas_comerciales",0)),
        "var_db":   pct(pc.get("deudas_bancarias",0), b.get("pasivo_corriente_anterior",{}).get("deudas_bancarias",0)),
        "estructura": "mejoró" if (pct(at,at_ant) or 0) > (pct(pt,pt_ant) or 0) else "empeoró",
    }

# ── Generación de texto ───────────────────────────────────────────────────────
def generar_texto(b, n, ix, api_key):
    prompt = f"""Sos un analista de crédito del Banco de la Nación Argentina.
Generá el texto de las tres secciones de un informe BCE en base a los datos numéricos provistos.

DATOS DEL BALANCE:
{json.dumps(b, ensure_ascii=False)}

DATOS DE NOSIS:
{json.dumps(n, ensure_ascii=False)}

ÍNDICES CALCULADOS:
{json.dumps(ix, ensure_ascii=False)}

INSTRUCCIONES:
- Estilo formal, técnico, tercera persona, párrafos corridos sin subtítulos internos
- Cifras en miles de pesos con signo $ y punto como separador (ej: $1.234.567)
- Porcentajes con un decimal (ej: 18.8%)
- Usá los índices exactamente como se proveen, no los recalcules
- Si una variación es null, no la menciones
- Cada sección: UN párrafo largo y continuo

CRITERIOS DE CALIFICACIÓN:
- Patrimonial: BUENA (solvencia >1.5), REGULAR (1.1-1.5), COMPROMETIDA (<1.1)
- Financiera: BUENA (liquidez corriente >1.5), REGULAR (1.0-1.5), COMPROMETIDA (<1.0)
- Económica: BUENA (margen neto >10%), REGULAR (3-10%), AJUSTADA (0-3%), INSUFICIENTE (<0)

Devolvé ÚNICAMENTE este JSON sin texto adicional ni backticks:
{{"patrimonial_calificacion":"","patrimonial_texto":"","financiera_calificacion":"","financiera_texto":"","economica_calificacion":"","economica_texto":""}}

CONTENIDO MÍNIMO:
PATRIMONIAL: solvencia, endeudamiento, % financiado con recursos propios, variaciones de activo corriente (con sus componentes principales), bienes de uso, pasivo corriente (con sus componentes), conclusión estructura, cobertura de activos fijos con PN, datos NOSIS (SRT, empleados, endeudamiento BCRA hoy vs cierre).
FINANCIERA: liquidez corriente, liquidez ácida y dependencia de bienes de cambio, Score NOSIS y situación BCRA, compromisos mensuales y % sobre ventas, cheques rechazados y deuda previsional, variación de liquidez vs ejercicio anterior.
ECONÓMICA: variación de ventas, margen bruto (actual vs anterior), promedio mensual de ventas, gastos comercialización y administración (% y variación), gastos financieros (%), EBITDA (monto y %), RECPAM (%), variación resultado operativo, resultado neto (margen actual vs anterior), rentabilidades (ventas, activo, PN).
"""
    return llamar_api_texto(prompt, api_key, max_tokens=3000)

# ── Generación del .docx ──────────────────────────────────────────────────────
def generar_docx(b, n, textos) -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import copy

    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin    = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin   = Cm(2)
        section.right_margin  = Cm(2)

    # Fuente por defecto
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)

    def add_par(text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.bold = bold
        return p

    def add_section_title(title, cal):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(4)
        r1 = p.add_run(f"{title} ({cal})")
        r1.font.name = 'Arial'; r1.font.size = Pt(10); r1.bold = True

    # ── Encabezado (tabla) ────────────────────────────────────────────────────
    razon_social = b.get("razon_social", "")
    cuit_raw     = b.get("cuit", "").replace("-", "")
    cuit = f"{cuit_raw[:2]}-{cuit_raw[2:10]}-{cuit_raw[10]}" if len(cuit_raw) == 11 else b.get("cuit","")

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    # Celda izquierda
    cell_left  = table.cell(0, 0)
    cell_right = table.cell(0, 1)

    # Shading gris en ambas celdas
    def set_cell_bg(cell, color_hex):
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), color_hex)
        tcPr.append(shd)

    set_cell_bg(cell_left,  'D9D9D9')
    set_cell_bg(cell_right, 'D9D9D9')

    # Contenido celda izquierda
    cell_left.text = ""
    p_left = cell_left.paragraphs[0]
    r = p_left.add_run("CASA N° – CENTRO COMERCIAL SL")
    r.font.name = 'Arial'; r.font.size = Pt(9); r.bold = True
    p2 = cell_left.add_paragraph("ZONAL SAN LUIS")
    p2.runs[0].font.name = 'Arial'; p2.runs[0].font.size = Pt(9); p2.runs[0].bold = True

    # Contenido celda derecha
    cell_right.text = ""
    pr = cell_right.paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = pr.add_run("CONCLUSIONES QUE SURGEN DEL ESTUDIO")
    r1.font.name = 'Arial'; r1.font.size = Pt(9); r1.bold = True
    p_rs = cell_right.add_paragraph(razon_social)
    p_rs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rs.runs[0].font.name = 'Arial'; p_rs.runs[0].font.size = Pt(11); p_rs.runs[0].bold = True
    p_cuit = cell_right.add_paragraph(cuit)
    p_cuit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cuit.runs[0].font.name = 'Arial'; p_cuit.runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # ── Intro ──────────────────────────────────────────────────────────────────
    domicilio    = b.get("domicilio", "")
    actividad    = b.get("actividad", "")
    fecha_cierre = b.get("fecha_cierre", "")
    nro_ej       = b.get("nro_ejercicio", "")
    consejo      = b.get("consejo_protocolizante", "")

    add_par(f"De los Estados Contables de la firma {razon_social}, C.U.I.T. {cuit} "
            f"con domicilio legal en {domicilio}. ACTIVIDAD ECONÓMICA: {actividad}.")

    doc.add_paragraph()

    # ── Consideraciones previas ────────────────────────────────────────────────
    add_par("CONSIDERACIONES PREVIAS", bold=True, space_after=2)
    add_par(f"La presente información se confecciona en base a las cifras en miles obtenidas "
            f"de los Estados contables correspondientes al ejercicio cerrado el {fecha_cierre} "
            f"N° {nro_ej}, auditados por Contador Público y protocolizados por el {consejo} "
            f"y demás información complementaria provista por la firma. Estimación realizada en miles.")

    # Socios
    socios = n.get("socios", [])
    if socios:
        doc.add_paragraph()
        add_par(f"{razon_social} está compuesta por:", space_after=2)
        for s in socios:
            nombre = s.get("nombre", "")
            pct_s  = s.get("porcentaje", "")
            if isinstance(pct_s, float):
                pct_s = f"{round(pct_s * 100)}%"
            add_par(f"- {nombre} ({pct_s})",
                    space_before=0, space_after=1)

    doc.add_paragraph()

    # ── Tres secciones ─────────────────────────────────────────────────────────
    for key_cal, key_txt, titulo in [
        ("patrimonial_calificacion", "patrimonial_texto", "SITUACIÓN PATRIMONIAL"),
        ("financiera_calificacion",  "financiera_texto",  "SITUACIÓN FINANCIERA"),
        ("economica_calificacion",   "economica_texto",   "SITUACIÓN ECONÓMICA"),
    ]:
        cal = textos.get(key_cal, "")
        txt = textos.get(key_txt, "")
        add_section_title(titulo, cal)
        for linea in txt.split("\n"):
            linea = linea.strip()
            if linea:
                add_par(linea)
        doc.add_paragraph()

    # Guardar en bytes
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp.name)
        tmp_path = tmp.name

    with open(tmp_path, "rb") as f:
        contenido = f.read()
    os.unlink(tmp_path)
    return contenido

# ── Pantalla principal ────────────────────────────────────────────────────────
def pantalla_principal():
    st.markdown("""
    <div class="header-bna">
        <div style="font-size:2rem">🏦</div>
        <div>
            <h1>Informe Automático BCE</h1>
            <p>Generación automática de informes de balance para evaluación crediticia</p>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Bienvenida y cierre de sesión
    col_usr, col_out = st.columns([4, 1])
    with col_usr:
        st.markdown(f"👤 **{st.session_state.get('usuario', '')}**")
    with col_out:
        if st.button("Salir", key="logout"):
            st.session_state.clear()
            st.rerun()

    # API Key
    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        st.markdown('<div class="info-box">⚠️ <strong>API Key no configurada.</strong> '
                    'Configurala como secret en Streamlit Cloud con el nombre '
                    '<code>ANTHROPIC_API_KEY</code>.</div>', unsafe_allow_html=True)
        api_key = st.text_input("O ingresá tu API Key aquí (temporaria):",
                                type="password", key="api_key_input")

    st.markdown("---")
    st.markdown("#### Subir documentos")
    st.markdown('<div class="info-box">📋 Subí el balance en PDF y el informe NOSIS. '
                'El sistema extrae los datos, calcula los índices y genera el informe Word listo para usar.</div>',
                unsafe_allow_html=True)

    # Upload de archivos
    col1, col2 = st.columns(2)
    with col1:
        st.markdown('<div class="upload-label">📄 Balance (EECC)</div>', unsafe_allow_html=True)
        archivo_balance = st.file_uploader("", type=["pdf"], key="balance",
                                           label_visibility="collapsed")
        if archivo_balance:
            st.success(f"✅ {archivo_balance.name}")
    with col2:
        st.markdown('<div class="upload-label">📊 Informe NOSIS</div>', unsafe_allow_html=True)
        archivo_nosis = st.file_uploader("", type=["pdf"], key="nosis",
                                         label_visibility="collapsed")
        if archivo_nosis:
            st.success(f"✅ {archivo_nosis.name}")

    st.markdown("")

    # Botón generar
    puede_generar = bool(archivo_balance and archivo_nosis and api_key)
    if st.button("⚡ Generar informe BCE", disabled=not puede_generar):
        _ejecutar_generacion(archivo_balance, archivo_nosis, api_key)


def _ejecutar_generacion(archivo_balance, archivo_nosis, api_key):
    """Ejecuta el pipeline completo mostrando progreso paso a paso."""

    pasos = st.empty()
    estado = []

    def mostrar(texto, tipo="activo"):
        cls = {"activo": "step-box", "done": "step-box done", "error": "step-box error"}[tipo]
        estado_html = "".join(
            f'<div class="{cls if i == len(estado)-1 else "step-box done"}">{e}</div>'
            for i, e in enumerate(estado)
        )
        pasos.markdown(estado_html, unsafe_allow_html=True)

    try:
        # Paso 1: Leer PDFs
        estado.append("📂 Leyendo PDFs...")
        mostrar("", "activo")
        bytes_balance = archivo_balance.read()
        bytes_nosis   = archivo_nosis.read()
        b64_balance   = pdf_a_base64(bytes_balance)
        b64_nosis     = pdf_a_base64(bytes_nosis)

        # Paso 2: Extraer balance
        estado[-1] = "✅ PDFs leídos"
        estado.append("🔍 Extrayendo datos del balance (puede tardar ~20s)...")
        mostrar("", "activo")
        datos_balance = llamar_api(PROMPT_BALANCE, b64_balance, 2000, api_key)

        # Paso 3: Extraer NOSIS
        estado[-1] = f"✅ Balance extraído: {datos_balance.get('razon_social','')}"
        estado.append("🔍 Extrayendo datos de NOSIS...")
        mostrar("", "activo")
        datos_nosis = llamar_api(PROMPT_NOSIS, b64_nosis, 1000, api_key)

        # Paso 4: Calcular índices
        estado[-1] = f"✅ NOSIS extraído: Score {datos_nosis.get('score','')}"
        estado.append("🔢 Calculando índices financieros...")
        mostrar("", "activo")
        indices = calcular_indices(datos_balance)

        # Paso 5: Generar texto
        estado[-1] = "✅ Índices calculados"
        estado.append("🤖 Generando texto con IA (puede tardar ~20s)...")
        mostrar("", "activo")
        textos = generar_texto(datos_balance, datos_nosis, indices, api_key)

        # Paso 6: Armar Word
        estado[-1] = f"✅ Texto generado"
        estado.append("📝 Armando documento Word...")
        mostrar("", "activo")
        docx_bytes = generar_docx(datos_balance, datos_nosis, textos)

        # Éxito
        estado[-1] = "✅ Documento generado"
        mostrar("", "done")

        razon = datos_balance.get("razon_social", "empresa")
        nro   = datos_balance.get("nro_ejercicio", "")
        nombre_archivo = f"Informe_BCE_{nro}_{razon.replace(' ', '_')}.docx"

        st.markdown(f"""
        <div class="result-card">
            <h3>✅ Informe generado correctamente</h3>
            <p>{razon} — Ejercicio N° {nro} — {datos_balance.get('fecha_cierre','')}</p>
        </div>
        """, unsafe_allow_html=True)

        st.download_button(
            label="⬇️ Descargar informe Word",
            data=docx_bytes,
            file_name=nombre_archivo,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    except Exception as e:
        estado.append(f"❌ Error: {str(e)}")
        mostrar("", "error")
        st.error(f"Ocurrió un error: {e}")


# ── Entry point ───────────────────────────────────────────────────────────────
if "autenticado" not in st.session_state:
    st.session_state["autenticado"] = False

if not st.session_state["autenticado"]:
    login()
else:
    pantalla_principal()
